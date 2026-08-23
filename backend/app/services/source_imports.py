import hashlib
import io
import re
from typing import Any

from sqlalchemy.orm import Session

from app.models import QuestionTypeEnum, SourceImport, SourceImportStatusEnum


def extract_source_segments(raw: bytes, content_type: str, filename: str) -> list[dict[str, Any]]:
    lower = filename.lower()
    if content_type == "application/pdf" or lower.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(raw))
        return [
            _segment(index + 1, f"Страница {index + 1}", page.extract_text() or "", "")
            for index, page in enumerate(reader.pages)
        ]
    if lower.endswith(".pptx") or content_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
        from pptx import Presentation

        presentation = Presentation(io.BytesIO(raw))
        segments: list[dict[str, Any]] = []
        for index, slide in enumerate(presentation.slides):
            texts = [shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
            title = texts[0] if texts else f"Слайд {index + 1}"
            notes = ""
            try:
                notes = slide.notes_slide.notes_text_frame.text.strip()
            except (AttributeError, KeyError):
                pass
            segments.append(_segment(index + 1, title, "\n".join(texts), notes))
        return segments
    raise ValueError("Only PPTX or PDF source files are supported")


def generate_source_candidates(db: Session, source: SourceImport) -> SourceImport:
    source.status = SourceImportStatusEnum.generating
    source.error_message = None
    db.add(source)
    db.commit()
    try:
        config = source.generation_config or {}
        count = max(1, min(int(config.get("count") or 8), 30))
        allowed = [str(item) for item in (config.get("question_types") or [QuestionTypeEnum.open_response.value])]
        excluded = set(source.excluded_segment_ids or [])
        usable = [segment for segment in source.segments if segment.get("id") not in excluded and _sentences(_segment_text(segment))]
        if not usable:
            raise ValueError("В выбранных слайдах не найден текст для генерации")
        facts = [(segment, sentence) for segment in usable for sentence in _sentences(_segment_text(segment))]
        candidates: list[dict[str, Any]] = []
        for index in range(count):
            segment, sentence = facts[index % len(facts)]
            requested_type = allowed[index % len(allowed)] if allowed else QuestionTypeEnum.open_response.value
            if requested_type not in {item.value for item in QuestionTypeEnum}:
                requested_type = QuestionTypeEnum.open_response.value
            candidates.append(_candidate(source.id, index, segment, sentence, requested_type, facts))
        source.candidates = candidates
        source.status = SourceImportStatusEnum.completed
    except Exception as exc:
        source.status = SourceImportStatusEnum.failed
        source.error_message = str(exc)[:4000]
    db.add(source)
    db.commit()
    db.refresh(source)
    return source


def _candidate(
    source_id: str,
    index: int,
    segment: dict[str, Any],
    sentence: str,
    question_type: str,
    facts: list[tuple[dict[str, Any], str]],
) -> dict[str, Any]:
    candidate_id = _stable_id(source_id, str(index), sentence)
    source_ref = {"source_import_id": source_id, "segment_id": segment["id"], "label": segment.get("title") or f"Слайд {segment['index']}"}
    base: dict[str, Any] = {
        "id": candidate_id,
        "status": "pending",
        "question_type": question_type,
        "expected_answer": sentence,
        "explanation": f"Ответ взят из раздела «{source_ref['label']}».",
        "source_refs": [source_ref],
        "answer_mode": "both",
        "max_score": 10,
        "competencies": [{"name": "Понимание материала", "weight": 1}],
        "options": [],
        "correct_option_ids": [],
    }
    if question_type == QuestionTypeEnum.open_response.value:
        base["text"] = f"Объясните своими словами: {sentence.rstrip('.!?')}?"
        return base

    distractors: list[str] = []
    for other_segment, fact in facts:
        if other_segment.get("id") != segment.get("id") and fact != sentence and fact not in distractors:
            distractors.append(fact)
        if len(distractors) == 3:
            break
    while len(distractors) < 3:
        distractors.append(f"В материале не утверждается: вариант {len(distractors) + 1}")
    correct = [_option(candidate_id, "correct-1", sentence)]
    if question_type == QuestionTypeEnum.multiple_choice.value:
        same_section = next((fact for other_segment, fact in facts if other_segment.get("id") == segment.get("id") and fact != sentence), None)
        second = same_section or f"Раздел раскрывает тему «{segment.get('title') or 'материал'}»"
        correct.append(_option(candidate_id, "correct-2", second))
        base["expected_answer"] = f"{sentence}\n{second}"
    options = correct + [_option(candidate_id, f"distractor-{i}", value) for i, value in enumerate(distractors, 1)]
    label = segment.get("title") or f"слайд {segment.get('index')}"
    base["text"] = f"Какое утверждение относится к разделу «{label}»?" if len(correct) == 1 else f"Выберите все утверждения, относящиеся к разделу «{label}»:"
    base["options"] = options
    base["correct_option_ids"] = [item["id"] for item in correct]
    base["answer_mode"] = "text"
    return base


def _segment(index: int, title: str, text: str, notes: str) -> dict[str, Any]:
    normalized = re.sub(r"\s+", " ", text).strip()
    normalized_notes = re.sub(r"\s+", " ", notes).strip()
    return {"id": f"segment-{index}", "index": index, "title": title[:240], "text": normalized, "notes": normalized_notes}


def _segment_text(segment: dict[str, Any]) -> str:
    return " ".join(part for part in [str(segment.get("text") or ""), str(segment.get("notes") or "")] if part).strip()


def _sentences(text: str) -> list[str]:
    values = [re.sub(r"\s+", " ", item).strip() for item in re.split(r"(?<=[.!?])\s+|\n+", text)]
    return [item[:600] for item in values if len(item) >= 25]


def _option(candidate_id: str, suffix: str, text: str) -> dict[str, str]:
    return {"id": _stable_id(candidate_id, suffix), "text": text[:1000]}


def _stable_id(*parts: str) -> str:
    return hashlib.sha256("|".join(parts).encode("utf-8")).hexdigest()[:20]
