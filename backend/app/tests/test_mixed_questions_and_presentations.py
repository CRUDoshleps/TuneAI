import io

from pptx import Presentation

from app.db.session import SessionLocal
from app.models import SourceImport
from app.services.source_imports import generate_source_candidates
from app.tests.conftest import auth_header, register_and_login


def test_choice_questions_unlock_in_order_and_score_without_ai(client):
    token = register_and_login(client, "admin@example.com")
    created = client.post(
        "/tests",
        headers=auth_header(token),
        json={
            "title": "Mixed knowledge check",
            "test_type": "exam",
            "questions": [
                {
                    "text": "Which pattern publishes database changes reliably?",
                    "question_type": "single_choice",
                    "options": [
                        {"id": "outbox", "text": "Transactional outbox"},
                        {"id": "cache", "text": "Browser cache"},
                    ],
                    "correct_option_ids": ["outbox"],
                    "explanation": "The outbox stores the event with the business transaction.",
                    "answer_mode": "text",
                    "order_index": 0,
                    "max_score": 4,
                },
                {
                    "text": "Select all retry-safe techniques.",
                    "question_type": "multiple_choice",
                    "options": [
                        {"id": "idempotency", "text": "Idempotency key"},
                        {"id": "dedupe", "text": "Consumer deduplication"},
                        {"id": "random", "text": "Random duplicate writes"},
                    ],
                    "correct_option_ids": ["idempotency", "dedupe"],
                    "answer_mode": "text",
                    "order_index": 1,
                    "max_score": 6,
                },
            ],
        },
    )
    assert created.status_code == 201, created.text
    test = created.json()
    client.patch(f"/tests/{test['id']}", headers=auth_header(token), json={"status": "published"})
    student_token = register_and_login(client, "student@example.com")
    student = client.get("/auth/me", headers=auth_header(student_token)).json()
    assigned = client.post(f"/tests/{test['id']}/assign", headers=auth_header(token), json={"user_id": student["id"]})
    assert assigned.status_code == 204, assigned.text
    attempt = client.post("/attempts", headers=auth_header(student_token), json={"test_id": test["id"]}).json()
    assert len(attempt["questions"]) == 1
    first = client.post(
        f"/attempts/{attempt['id']}/questions/{test['questions'][0]['id']}/choices",
        headers={**auth_header(student_token), "Idempotency-Key": "choice-first"},
        json={"selected_option_ids": ["outbox"]},
    )
    assert first.status_code == 201, first.text
    assert first.json()["status"] == "processing"
    assert len(first.json()["questions"]) == 2
    second = client.post(
        f"/attempts/{attempt['id']}/questions/{test['questions'][1]['id']}/choices",
        headers={**auth_header(student_token), "Idempotency-Key": "choice-second"},
        json={"selected_option_ids": ["idempotency", "dedupe"]},
    )
    assert second.status_code == 201, second.text
    assert second.json()["status"] == "completed"
    assert second.json()["total_score"] == 10
    assert second.json()["answers"][0]["evaluation"]["evaluation_version"] == "deterministic-choice-v1"


def test_pptx_import_generates_mixed_candidates_and_accepts_question(client):
    token = register_and_login(client, "admin@example.com")
    test = client.post(
        "/tests",
        headers=auth_header(token),
        json={"title": "Presentation assessment", "test_type": "exam", "questions": [{"text": "Starter question", "expected_answer": "Starter"}]},
    ).json()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Transactional outbox"
    slide.placeholders[1].text = "Transactional outbox writes the domain change and outgoing event in one database transaction."
    second = presentation.slides.add_slide(presentation.slide_layouts[1])
    second.shapes.title.text = "Idempotency"
    second.placeholders[1].text = "An idempotency key prevents repeated requests from creating duplicate business effects."
    content = io.BytesIO()
    presentation.save(content)
    uploaded = client.post(
        f"/source-imports/upload?test_id={test['id']}",
        headers=auth_header(token),
        files={"file": ("lecture.pptx", content.getvalue(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
    )
    assert uploaded.status_code == 201, uploaded.text
    source_payload = uploaded.json()
    assert source_payload["status"] == "ready"
    assert len(source_payload["segments"]) == 2

    with SessionLocal() as db:
        source = db.get(SourceImport, source_payload["id"])
        source.generation_config = {"count": 3, "question_types": ["open_response", "single_choice", "multiple_choice"]}
        db.add(source)
        db.commit()
        generated = generate_source_candidates(db, source)
        candidate_id = generated.candidates[1]["id"]
        assert {item["question_type"] for item in generated.candidates} == {"open_response", "single_choice", "multiple_choice"}

    accepted = client.post(
        f"/source-imports/{source_payload['id']}/candidates/{candidate_id}/accept",
        headers=auth_header(token),
    )
    assert accepted.status_code == 201, accepted.text
    assert accepted.json()["question_type"] == "single_choice"
    assert accepted.json()["correct_option_ids"]
    source_after_accept = client.get(
        f"/source-imports/{source_payload['id']}",
        headers=auth_header(token),
    )
    assert source_after_accept.status_code == 200, source_after_accept.text
    accepted_candidate = next(
        item for item in source_after_accept.json()["candidates"] if item["id"] == candidate_id
    )
    assert accepted_candidate["status"] == "accepted"
    assert accepted_candidate["question_id"] == accepted.json()["id"]
    accepted_again = client.post(
        f"/source-imports/{source_payload['id']}/candidates/{candidate_id}/accept",
        headers=auth_header(token),
    )
    assert accepted_again.status_code == 201, accepted_again.text
    assert accepted_again.json()["id"] == accepted.json()["id"]
    test_after_accept = client.get(f"/tests/{test['id']}", headers=auth_header(token))
    assert test_after_accept.status_code == 200, test_after_accept.text
    assert len(test_after_accept.json()["questions"]) == 2
    audit = client.get("/admin/audit-log", headers=auth_header(token))
    assert audit.status_code == 200
    assert any(item["action"] == "question.accept_generated" for item in audit.json())
